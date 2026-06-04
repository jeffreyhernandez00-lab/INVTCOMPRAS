from io import BytesIO
import re
import unicodedata

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl.styles import Font, PatternFill


REQUIRED_COLUMNS = [
    "Código de producto",
    "Nombre de producto",
    "Laboratorio",
    "Venta mes 1",
    "Venta mes 2",
    "Venta mes 3",
    "Venta mes 4",
    "Venta mes 5",
    "Venta mes 6",
    "Costo unitario",
    "Inventario actual",
]

SALES_COLUMNS = [
    "Venta mes 1",
    "Venta mes 2",
    "Venta mes 3",
    "Venta mes 4",
    "Venta mes 5",
    "Venta mes 6",
]

NUMERIC_COLUMNS = [
    *SALES_COLUMNS,
    "Costo unitario",
    "Inventario actual",
]

OUTPUT_COLUMNS = [
    "Código",
    "Producto",
    "Laboratorio",
    "Responsable",
    "Venta mes 1",
    "Venta mes 2",
    "Venta mes 3",
    "Venta mes 4",
    "Venta mes 5",
    "Venta mes 6",
    "Consumo 6 meses",
    "Promedio mensual",
    "Promedio diario",
    "Costo unitario",
    "Total costo",
    "% participación",
    "% participación acumulado",
    "Clasificación ABC",
    "Desviación estándar",
    "Rango",
    "Inventario actual",
    "Punto de reorden",
    "Inventario mínimo",
    "Inventario máximo",
    "Diferencia vs máximo",
    "Días de consumo disponibles",
    "Sugerido de compra",
    "Estado inventario",
]

PURCHASE_COLUMNS = [
    "Código",
    "Producto",
    "Laboratorio",
    "Responsable",
    "Clasificación ABC",
    "Inventario actual",
    "Punto de reorden",
    "Inventario máximo",
    "Sugerido de compra",
    "Costo unitario",
    "Total estimado compra",
]

RESPONSIBLE_SHEET_NAME = "Responsables laboratorio"
RESPONSIBLE_COLUMNS = ["Laboratorio", "Responsable"]


def set_page_style() -> None:
    st.markdown(
        """
        <style>
            :root {
                --brand-orange: #d27a00;
                --brand-green: #70b62c;
                --brand-dark-green: #007a45;
                --brand-ink: #1f2937;
                --brand-muted: #667085;
                --brand-border: #d9e4d0;
                --brand-surface: #ffffff;
                --brand-wash: #f4f8ef;
            }
            .stApp {
                background:
                    linear-gradient(180deg, rgba(112, 182, 44, 0.10), rgba(255, 255, 255, 0) 260px),
                    #f7f9f3;
                color: var(--brand-ink);
            }
            .stApp h1,
            .stApp h2,
            .stApp h3,
            .stApp p,
            .stApp label,
            .stApp [data-testid="stMarkdownContainer"] {
                color: var(--brand-ink);
            }
            .stApp h2,
            .stApp h3 {
                border-left: 5px solid var(--brand-orange);
                padding-left: 10px;
            }
            div[data-testid="stMetric"] {
                background: var(--brand-surface);
                border: 1px solid var(--brand-border);
                border-radius: 8px;
                padding: 12px 14px;
                box-shadow: 0 8px 24px rgba(20, 83, 45, 0.08);
            }
            div[data-testid="stDataFrame"] {
                border: 1px solid var(--brand-border);
                border-radius: 8px;
                overflow: hidden;
                box-shadow: 0 8px 24px rgba(20, 83, 45, 0.06);
            }
            .section-note {
                color: var(--brand-muted);
                font-size: 0.94rem;
                margin-bottom: 0.65rem;
            }
            .formula-note {
                background: #ffffff;
                border: 1px solid var(--brand-border);
                border-left: 5px solid var(--brand-green);
                border-radius: 8px;
                color: #344054;
                padding: 12px 14px;
                margin: 8px 0 16px;
                box-shadow: 0 6px 18px rgba(20, 83, 45, 0.06);
            }
            .brand-header {
                align-items: center;
                background:
                    linear-gradient(135deg, rgba(0, 122, 69, 0.96), rgba(112, 182, 44, 0.92)),
                    #007a45;
                border: 1px solid rgba(255, 255, 255, 0.36);
                border-radius: 8px;
                box-shadow: 0 18px 42px rgba(0, 92, 52, 0.20);
                display: flex;
                gap: 18px;
                margin: 8px 0 22px;
                padding: 22px 24px;
            }
            .brand-logo {
                align-items: center;
                background: #f8faf7;
                border: 5px solid var(--brand-orange);
                border-radius: 50%;
                display: flex;
                flex: 0 0 84px;
                height: 84px;
                justify-content: center;
                position: relative;
                width: 84px;
            }
            .brand-logo::before {
                background: var(--brand-dark-green);
                border-radius: 8px;
                content: "";
                height: 58px;
                position: absolute;
                transform: rotate(5deg);
                width: 28px;
            }
            .brand-logo::after {
                background: var(--brand-green);
                border-radius: 8px;
                content: "";
                height: 30px;
                position: absolute;
                width: 70px;
            }
            .brand-copy h1 {
                color: #ffffff;
                font-size: 2rem;
                line-height: 1.1;
                margin: 0 0 6px;
            }
            .brand-copy p {
                color: rgba(255, 255, 255, 0.90);
                font-size: 1rem;
                margin: 0;
            }
            .brand-pill-row {
                display: flex;
                flex-wrap: wrap;
                gap: 8px;
                margin-top: 12px;
            }
            .brand-pill {
                background: rgba(255, 255, 255, 0.14);
                border: 1px solid rgba(255, 255, 255, 0.28);
                border-radius: 999px;
                color: #ffffff;
                font-size: 0.86rem;
                padding: 5px 10px;
            }
            div.stButton > button,
            div.stDownloadButton > button {
                border-radius: 8px;
                border: 1px solid var(--brand-dark-green);
                font-weight: 700;
            }
            div.stButton > button[kind="primary"] {
                background: var(--brand-orange);
                border-color: var(--brand-orange);
                color: #ffffff;
            }
            div.stDownloadButton > button {
                background: #ffffff;
                color: var(--brand-dark-green);
            }
            div.stDownloadButton > button:hover {
                border-color: var(--brand-orange);
                color: var(--brand-orange);
            }
            div[data-baseweb="input"] {
                border-radius: 8px;
            }
            @media (max-width: 720px) {
                .brand-header {
                    align-items: flex-start;
                    flex-direction: column;
                    padding: 18px;
                }
                .brand-copy h1 {
                    font-size: 1.55rem;
                }
            }
        </style>
        """,
        unsafe_allow_html=True,
    )


def render_brand_header() -> None:
    st.markdown(
        """
        <div class="brand-header">
            <div class="brand-logo" aria-hidden="true"></div>
            <div class="brand-copy">
                <h1>Sugerido de compra automático</h1>
                <p>Analiza ventas, inventario actual y clasificación ABC para decidir qué comprar y cuánto pedir.</p>
                <div class="brand-pill-row">
                    <span class="brand-pill">Inventario</span>
                    <span class="brand-pill">ABC</span>
                    <span class="brand-pill">Reorden</span>
                    <span class="brand-pill">Exportación Excel</span>
                </div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def validate_columns(data: pd.DataFrame) -> list[str]:
    available_columns = {normalize_text(column) for column in data.columns}
    return [
        column
        for column in REQUIRED_COLUMNS
        if normalize_text(column) not in available_columns
    ]


def normalize_text(value: object) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"\s+", " ", text)


def canonicalize_columns(data: pd.DataFrame) -> pd.DataFrame:
    column_lookup = {normalize_text(column): column for column in data.columns}
    canonical_data = pd.DataFrame(index=data.index)

    for required_column in REQUIRED_COLUMNS:
        source_column = column_lookup[normalize_text(required_column)]
        canonical_data[required_column] = data[source_column]

    return canonical_data


def read_responsible_sheet(excel_file: pd.ExcelFile) -> pd.DataFrame:
    sheet_lookup = {normalize_text(sheet_name): sheet_name for sheet_name in excel_file.sheet_names}
    sheet_name = sheet_lookup.get(normalize_text(RESPONSIBLE_SHEET_NAME))

    if sheet_name is None:
        return pd.DataFrame(columns=RESPONSIBLE_COLUMNS)

    responsible_data = pd.read_excel(excel_file, sheet_name=sheet_name)
    column_lookup = {normalize_text(column): column for column in responsible_data.columns}

    if "laboratorio" not in column_lookup or "responsable" not in column_lookup:
        return pd.DataFrame(columns=RESPONSIBLE_COLUMNS)

    cleaned = pd.DataFrame(
        {
            "Laboratorio": responsible_data[column_lookup["laboratorio"]],
            "Responsable": responsible_data[column_lookup["responsable"]],
        }
    )
    cleaned["Laboratorio"] = cleaned["Laboratorio"].astype(str).str.strip()
    cleaned["Responsable"] = cleaned["Responsable"].fillna("").astype(str).str.strip()
    cleaned = cleaned[cleaned["Laboratorio"].ne("")]
    return cleaned.drop_duplicates(subset=["Laboratorio"], keep="last")


def render_responsible_editor(
    input_data: pd.DataFrame,
    responsible_data: pd.DataFrame,
    editor_key: str,
) -> dict[str, str]:
    product_data = canonicalize_columns(input_data)
    laboratories = (
        product_data["Laboratorio"]
        .fillna("")
        .astype(str)
        .str.strip()
        .loc[lambda values: values.ne("")]
        .drop_duplicates()
        .sort_values()
        .tolist()
    )
    responsible_lookup = dict(
        zip(responsible_data["Laboratorio"], responsible_data["Responsable"])
    )
    editor_data = pd.DataFrame(
        {
            "Laboratorio": laboratories,
            "Responsable": [
                responsible_lookup.get(laboratory, "") for laboratory in laboratories
            ],
        }
    )

    st.subheader("Responsables por laboratorio")
    st.markdown(
        """
        <div class="formula-note">
            Define quién será responsable de cada laboratorio. Esta información se agregará a los resultados
            y a los Excel exportados.
        </div>
        """,
        unsafe_allow_html=True,
    )
    edited_data = st.data_editor(
        editor_data,
        hide_index=True,
        num_rows="fixed",
        use_container_width=True,
        disabled=["Laboratorio"],
        column_config={
            "Laboratorio": st.column_config.TextColumn("Laboratorio"),
            "Responsable": st.column_config.TextColumn(
                "Responsable",
                help="Persona encargada de revisar o gestionar la compra de este laboratorio.",
            ),
        },
        key=editor_key,
    )

    return dict(zip(edited_data["Laboratorio"], edited_data["Responsable"]))


@st.cache_data(show_spinner=False)
def build_template() -> bytes:
    product_template = pd.DataFrame(
        [
            {
                "Código de producto": "SKU-001",
                "Nombre de producto": "Producto ejemplo",
                "Laboratorio": "Proveedor ejemplo",
                "Venta mes 1": 120,
                "Venta mes 2": 135,
                "Venta mes 3": 128,
                "Venta mes 4": 142,
                "Venta mes 5": 138,
                "Venta mes 6": 150,
                "Costo unitario": 12.5,
                "Inventario actual": 40,
            }
        ],
        columns=REQUIRED_COLUMNS,
    )
    abc_template = pd.DataFrame(
        [
            {
                "Clasificación ABC": "A",
                "Días de inventario": 30,
                "Entrega proveedor (días)": 15,
                "Buffer seguridad (días)": 10,
            },
            {
                "Clasificación ABC": "B",
                "Días de inventario": 25,
                "Entrega proveedor (días)": 10,
                "Buffer seguridad (días)": 5,
            },
            {
                "Clasificación ABC": "C",
                "Días de inventario": 20,
                "Entrega proveedor (días)": 7,
                "Buffer seguridad (días)": 3,
            },
        ]
    )
    config_template = pd.DataFrame(
        [
            {
                "Parámetro": "Días hábiles de consumo por mes",
                "Valor sugerido": 26,
                "Descripción": "Se usa para calcular Promedio diario = Promedio mensual / días hábiles.",
            }
        ]
    )
    responsible_template = pd.DataFrame(
        [
            {
                "Laboratorio": "Proveedor ejemplo",
                "Responsable": "Nombre del responsable",
            }
        ]
    )
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        product_template.to_excel(writer, index=False, sheet_name="Productos")
        responsible_template.to_excel(
            writer,
            index=False,
            sheet_name="Responsables laboratorio",
        )
        abc_template.to_excel(writer, index=False, sheet_name="Parámetros ABC")
        config_template.to_excel(writer, index=False, sheet_name="Configuración")

        for worksheet in writer.sheets.values():
            worksheet.freeze_panes = "A2"
            for column_cells in worksheet.columns:
                max_length = max(len(str(cell.value or "")) for cell in column_cells)
                worksheet.column_dimensions[column_cells[0].column_letter].width = min(
                    max_length + 3,
                    55,
                )
    return buffer.getvalue()


def to_excel_bytes(
    data: pd.DataFrame,
    sheet_name: str,
    percentage_columns: list[str] | None = None,
) -> bytes:
    percentage_columns = percentage_columns or []
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        data.to_excel(writer, index=False, sheet_name=sheet_name)
        worksheet = writer.sheets[sheet_name]
        worksheet.freeze_panes = "A2"
        header_fill = PatternFill("solid", fgColor="007A45")
        header_font = Font(color="FFFFFF", bold=True)
        excess_fill = PatternFill("solid", fgColor="FCE4D6")
        purchase_fill = PatternFill("solid", fgColor="E2F0D9")

        for header_cell in worksheet[1]:
            header_cell.fill = header_fill
            header_cell.font = header_font

        for column_index, column_name in enumerate(data.columns, start=1):
            column_letter = worksheet.cell(row=1, column=column_index).column_letter
            max_length = max(
                len(str(column_name)),
                *(len(str(value)) for value in data[column_name].head(200).fillna("")),
            )
            worksheet.column_dimensions[column_letter].width = min(max_length + 3, 38)

            if column_name in percentage_columns:
                for cell in worksheet[column_letter][1:]:
                    cell.number_format = "0.00%"

            if column_name in {
                "Costo unitario",
                "Total costo",
                "Total estimado compra",
                "Diferencia vs máximo",
            }:
                for cell in worksheet[column_letter][1:]:
                    cell.number_format = '#,##0.00'

        if "Estado inventario" in data.columns:
            state_column = data.columns.get_loc("Estado inventario") + 1
            for row_index in range(2, len(data) + 2):
                state = worksheet.cell(row=row_index, column=state_column).value
                if state == "Exceso":
                    fill = excess_fill
                elif state == "Compra sugerida":
                    fill = purchase_fill
                else:
                    continue
                for cell in worksheet[row_index]:
                    cell.fill = fill
    return buffer.getvalue()


def clean_input_data(data: pd.DataFrame) -> pd.DataFrame:
    cleaned = canonicalize_columns(data)
    cleaned = cleaned.dropna(how="all")

    for column in NUMERIC_COLUMNS:
        cleaned[column] = pd.to_numeric(cleaned[column], errors="coerce").fillna(0)

    cleaned["Código de producto"] = cleaned["Código de producto"].astype(str).str.strip()
    cleaned["Nombre de producto"] = cleaned["Nombre de producto"].astype(str).str.strip()
    cleaned["Laboratorio"] = cleaned["Laboratorio"].astype(str).str.strip()
    return cleaned


def classify_products(
    accumulated_participation: pd.Series,
    a_limit: float,
    b_limit: float,
) -> pd.Series:
    conditions = [
        accumulated_participation <= a_limit,
        accumulated_participation <= b_limit,
    ]
    return pd.Series(
        np.select(conditions, ["A", "B"], default="C"),
        index=accumulated_participation.index,
    )


def calculate_purchase_suggestion(
    data: pd.DataFrame,
    abc_parameters: dict[str, dict[str, float]],
    a_limit: float,
    b_limit: float,
    working_days: float,
    lab_responsibles: dict[str, str] | None = None,
) -> pd.DataFrame:
    result = clean_input_data(data)
    lab_responsibles = lab_responsibles or {}
    result["Responsable"] = result["Laboratorio"].map(lab_responsibles).fillna("")

    result["Consumo 6 meses"] = result[SALES_COLUMNS].sum(axis=1)
    result["Promedio mensual"] = result["Consumo 6 meses"] / 6
    result["Promedio diario"] = result["Promedio mensual"] / working_days
    result["Total costo"] = result["Consumo 6 meses"] * result["Costo unitario"]
    result["Desviación estándar"] = result[SALES_COLUMNS].std(axis=1, ddof=0)
    result["Rango"] = result[SALES_COLUMNS].max(axis=1) - result[SALES_COLUMNS].min(axis=1)

    result = result.sort_values("Total costo", ascending=False).reset_index(drop=True)
    total_general = result["Total costo"].sum()

    if total_general > 0:
        result["% participación"] = result["Total costo"] / total_general
    else:
        result["% participación"] = 0

    result["% participación acumulado"] = result["% participación"].cumsum()
    result["Clasificación ABC"] = classify_products(
        result["% participación acumulado"] * 100,
        a_limit,
        b_limit,
    )

    max_days = result["Clasificación ABC"].map(
        lambda classification: abc_parameters[classification]["max_days"]
    )
    lead_time = result["Clasificación ABC"].map(
        lambda classification: abc_parameters[classification]["lead_time"]
    )
    buffer_days = result["Clasificación ABC"].map(
        lambda classification: abc_parameters[classification]["buffer_days"]
    )

    result["Inventario mínimo"] = result["Promedio diario"] * buffer_days
    result["Punto de reorden"] = result["Promedio diario"] * (lead_time + buffer_days)
    result["Inventario máximo"] = result["Promedio diario"] * (
        lead_time + buffer_days + max_days
    )
    result["Días de consumo disponibles"] = np.where(
        result["Promedio diario"] > 0,
        result["Inventario actual"] / result["Promedio diario"],
        np.inf,
    )

    purchase_raw = result["Inventario máximo"] - result["Inventario actual"]
    result["Diferencia vs máximo"] = purchase_raw
    applies_purchase = result["Inventario actual"] <= result["Punto de reorden"]
    has_excess = result["Inventario actual"] > result["Inventario máximo"]
    result["Sugerido de compra"] = np.where(
        applies_purchase,
        np.ceil(np.maximum(purchase_raw, 0)),
        0,
    ).astype(int)
    result["Estado inventario"] = np.select(
        [has_excess, result["Sugerido de compra"] > 0],
        ["Exceso", "Compra sugerida"],
        default="Normal",
    )
    result["Total estimado compra"] = (
        result["Sugerido de compra"] * result["Costo unitario"]
    )

    result = result.rename(
        columns={
            "Código de producto": "Código",
            "Nombre de producto": "Producto",
        }
    )

    return result[OUTPUT_COLUMNS + ["Total estimado compra"]]


def format_results_for_screen(data: pd.DataFrame) -> pd.DataFrame:
    formatted = data.copy()
    formatted["% participación"] = formatted["% participación"] * 100
    formatted["% participación acumulado"] = formatted["% participación acumulado"] * 100
    formatted["Días de consumo disponibles"] = formatted[
        "Días de consumo disponibles"
    ].replace(np.inf, 0)
    return formatted


def style_inventory_status(data: pd.DataFrame):
    def style_row(row: pd.Series) -> list[str]:
        if row["Estado inventario"] == "Exceso":
            return ["background-color: #fde7dc; color: #8a2c0d"] * len(row)
        if row["Estado inventario"] == "Compra sugerida":
            return ["background-color: #e7f4df; color: #1f5f2d"] * len(row)
        return [""] * len(row)

    return data.style.apply(style_row, axis=1)


def get_purchase_rows(data: pd.DataFrame) -> pd.DataFrame:
    return data.loc[data["Sugerido de compra"] > 0, PURCHASE_COLUMNS].copy()


def render_abc_parameters() -> tuple[dict[str, dict[str, float]], float, float]:
    st.subheader("Parámetros de inventario por clasificación ABC")
    st.markdown(
        """
        <div class="formula-note">
            <strong>Días de inventario:</strong> cobertura adicional que quieres tener en bodega.<br>
            <strong>Entrega proveedor (días):</strong> días que tarda el proveedor en entregar después del pedido.<br>
            <strong>Buffer seguridad (días):</strong> colchón adicional para evitar quiebres de inventario.
            <br><br>
            <strong>Fórmula aplicada:</strong> el inventario mínimo usa solo el buffer de seguridad;
            el punto de reorden usa entrega del proveedor más buffer, para comprar antes de llegar al mínimo.
            El inventario máximo suma entrega del proveedor, buffer y días de inventario.
        </div>
        """,
        unsafe_allow_html=True,
    )

    default_parameters = pd.DataFrame(
        [
            {
                "Clasificación ABC": "A",
                "Días de inventario": 30.0,
                "Entrega proveedor (días)": 15.0,
                "Buffer seguridad (días)": 10.0,
            },
            {
                "Clasificación ABC": "B",
                "Días de inventario": 25.0,
                "Entrega proveedor (días)": 10.0,
                "Buffer seguridad (días)": 5.0,
            },
            {
                "Clasificación ABC": "C",
                "Días de inventario": 20.0,
                "Entrega proveedor (días)": 7.0,
                "Buffer seguridad (días)": 3.0,
            },
        ]
    )

    edited_parameters = st.data_editor(
        default_parameters,
        hide_index=True,
        num_rows="fixed",
        use_container_width=True,
        disabled=["Clasificación ABC"],
        column_config={
            "Clasificación ABC": st.column_config.TextColumn(
                "Clasificación ABC",
                help="Grupo asignado por participación acumulada de costo.",
            ),
            "Días de inventario": st.column_config.NumberColumn(
                "Días de inventario",
                min_value=0.0,
                step=1.0,
                help="Días adicionales de consumo que deseas tener después de cubrir entrega y seguridad.",
            ),
            "Entrega proveedor (días)": st.column_config.NumberColumn(
                "Entrega proveedor (días)",
                min_value=0.0,
                step=1.0,
                help="Días que tarda el proveedor en entregar el pedido.",
            ),
            "Buffer seguridad (días)": st.column_config.NumberColumn(
                "Buffer seguridad (días)",
                min_value=0.0,
                step=1.0,
                help="Días extra de protección ante variaciones de venta o entrega.",
            ),
        },
        key="abc_parameter_table",
    )

    abc_parameters: dict[str, dict[str, float]] = {}
    for row in edited_parameters.to_dict("records"):
        classification = str(row["Clasificación ABC"])
        abc_parameters[classification] = {
            "max_days": float(row["Días de inventario"]),
            "lead_time": float(row["Entrega proveedor (días)"]),
            "buffer_days": float(row["Buffer seguridad (días)"]),
        }

    st.subheader("Límites de clasificación ABC")
    st.markdown(
        """
        <div class="formula-note">
            La clasificación se calcula con el porcentaje acumulado de participación:
            A llega hasta el primer límite, B hasta el segundo y C queda para lo que supere B.
        </div>
        """,
        unsafe_allow_html=True,
    )
    threshold_columns = st.columns(2)
    a_limit = threshold_columns[0].number_input(
        "A: hasta este % acumulado",
        min_value=0.0,
        max_value=100.0,
        value=80.0,
        step=1.0,
        help="Ejemplo: 80 significa que los productos hasta 80% acumulado quedan como A.",
    )
    b_limit = threshold_columns[1].number_input(
        "B: hasta este % acumulado",
        min_value=0.0,
        max_value=100.0,
        value=95.0,
        step=1.0,
        help="Ejemplo: 95 significa que lo mayor a 80% y hasta 95% queda como B.",
    )
    return abc_parameters, a_limit, b_limit


def render_summary(results: pd.DataFrame) -> None:
    purchase_rows = get_purchase_rows(results)
    counts = results["Clasificación ABC"].value_counts()
    excess_count = int((results["Estado inventario"] == "Exceso").sum())

    metric_columns = st.columns(5)
    metric_columns[0].metric("Productos analizados", f"{len(results):,}")
    metric_columns[1].metric(
        "Productos A / B / C",
        f"{counts.get('A', 0)} / {counts.get('B', 0)} / {counts.get('C', 0)}",
    )
    metric_columns[2].metric("Aplican compra", f"{len(purchase_rows):,}")
    metric_columns[3].metric("Con exceso", f"{excess_count:,}")
    metric_columns[4].metric(
        "Monto estimado compra",
        f"{purchase_rows['Total estimado compra'].sum():,.2f}",
    )


def main() -> None:
    st.set_page_config(
        page_title="Sugerido de Compra",
        layout="wide",
    )
    set_page_style()

    render_brand_header()

    st.subheader("Días de consumo")
    st.markdown(
        """
        <div class="formula-note">
            Este valor indica cuántos días hábiles de venta o consumo tiene un mes.
            Se usa para calcular: <strong>Promedio diario = Promedio mensual / días hábiles</strong>.
        </div>
        """,
        unsafe_allow_html=True,
    )
    working_days = st.number_input(
        "Días hábiles de consumo por mes",
        min_value=1.0,
        max_value=31.0,
        value=26.0,
        step=1.0,
        help="Ejemplo: usa 26 si normalmente vendes de lunes a sábado; usa 22 si solo consideras lunes a viernes.",
    )

    abc_parameters, a_limit, b_limit = render_abc_parameters()

    if a_limit >= b_limit:
        st.error("El límite de A debe ser menor que el límite de B.")
        st.stop()

    action_columns = st.columns([1, 1, 3])
    action_columns[0].download_button(
        "Descargar plantilla",
        data=build_template(),
        file_name="plantilla_sugerido_compra.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )

    uploaded_file = st.file_uploader(
        "Subir archivo Excel",
        type=["xlsx", "xls"],
        accept_multiple_files=False,
    )

    if uploaded_file is None:
        st.info("Sube un archivo Excel para ver la vista previa y procesar los datos.")
        return

    current_file_key = f"{uploaded_file.name}-{uploaded_file.size}"
    if st.session_state.get("uploaded_file_key") != current_file_key:
        st.session_state["uploaded_file_key"] = current_file_key
        st.session_state.pop("results", None)

    try:
        excel_file = pd.ExcelFile(uploaded_file)
        input_data = pd.read_excel(excel_file, sheet_name=0)
        responsible_data = read_responsible_sheet(excel_file)
    except Exception as error:
        st.error(f"No se pudo leer el archivo Excel: {error}")
        return

    missing_columns = validate_columns(input_data)
    if missing_columns:
        st.error("Faltan columnas obligatorias en el Excel.")
        st.write(pd.DataFrame({"Columnas faltantes": missing_columns}))
        return

    st.subheader("Vista previa")
    st.dataframe(input_data.head(50), use_container_width=True, hide_index=True)

    lab_responsibles = render_responsible_editor(
        input_data,
        responsible_data,
        f"responsible_editor_{current_file_key}",
    )

    calculate = st.button("Calcular sugerido", type="primary", use_container_width=False)
    if not calculate and "results" not in st.session_state:
        return

    if calculate:
        st.session_state["results"] = calculate_purchase_suggestion(
            input_data,
            abc_parameters,
            a_limit,
            b_limit,
            working_days,
            lab_responsibles,
        )

    results = st.session_state["results"]
    screen_results = format_results_for_screen(results)
    purchase_results = get_purchase_rows(results)

    st.subheader("Resumen general")
    render_summary(results)

    st.subheader("Resultados completos")
    st.dataframe(
        style_inventory_status(screen_results[OUTPUT_COLUMNS]),
        use_container_width=True,
        hide_index=True,
        column_config={
            "% participación": st.column_config.NumberColumn(format="%.2f %%"),
            "% participación acumulado": st.column_config.NumberColumn(format="%.2f %%"),
            "Costo unitario": st.column_config.NumberColumn(format="%.2f"),
            "Total costo": st.column_config.NumberColumn(format="%.2f"),
            "Promedio mensual": st.column_config.NumberColumn(format="%.2f"),
            "Promedio diario": st.column_config.NumberColumn(format="%.2f"),
            "Punto de reorden": st.column_config.NumberColumn(format="%.2f"),
            "Inventario mínimo": st.column_config.NumberColumn(format="%.2f"),
            "Inventario máximo": st.column_config.NumberColumn(format="%.2f"),
            "Diferencia vs máximo": st.column_config.NumberColumn(format="%.2f"),
            "Días de consumo disponibles": st.column_config.NumberColumn(format="%.2f"),
        },
    )

    export_columns = st.columns(2)
    export_columns[0].download_button(
        "Exportar Excel completo",
        data=to_excel_bytes(
            results[OUTPUT_COLUMNS],
            "Resultados completos",
            ["% participación", "% participación acumulado"],
        ),
        file_name="sugerido_compra_completo.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )
    export_columns[1].download_button(
        "Exportar compra sugerida",
        data=to_excel_bytes(purchase_results, "Compra sugerida"),
        file_name="compra_sugerida.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )


if __name__ == "__main__":
    main()
